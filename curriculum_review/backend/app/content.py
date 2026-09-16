"""Safe, normalized document import, comparison, anchoring and export."""

import base64
import hashlib
import io
import json
import re
import zipfile
from difflib import SequenceMatcher
from pathlib import Path
from uuid import uuid4
from docx import Document as WordDocument
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.image.image import Image as WordImage
from docx.oxml import OxmlElement
from docx.enum.style import WD_STYLE_TYPE
from docx.table import Table
from docx.text.paragraph import Paragraph
from fastapi import HTTPException
from lxml import etree
from markdown_it import MarkdownIt
from pypdf import PdfReader

MAX_UPLOAD = 20 * 1024 * 1024
MAX_TEXT = 200000
MAX_MBZ_UPLOAD = 1024 * 1024 * 1024


def block(kind, text="", **extra):
    return {"id": str(uuid4()), "type": kind, "text": text, **extra}


def text_of(content):
    return "\n\n".join(
        (
            "\n".join("\t".join(row) for row in b["rows"])
            if b["type"] == "table"
            else b.get("text", "")
        )
        for b in content["blocks"]
    )


def semantic_block(b):
    fields = {
        'paragraph': ('type','text'),
        'heading': ('type','text','level'),
        'list': ('type','text','level','ordered'),
        'table': ('type','rows'),
        'image': ('type','text','data','mime'),
    }[b['type']]
    return {key: b.get(key) for key in fields}


def fingerprint(content):
    clean = [semantic_block(b) for b in content['blocks']]
    return hashlib.sha256(json.dumps(clean, sort_keys=True).encode()).hexdigest()


def markdown(text):
    tokens = MarkdownIt("commonmark").enable("table").parse(text)
    blocks, heading, list_depth, ordered, rows, row, in_table = (
        [],
        0,
        0,
        False,
        [],
        [],
        False,
    )
    for token in tokens:
        if token.type == "heading_open":
            heading = int(token.tag[1:])
        elif token.type == "heading_close":
            heading = 0
        elif token.type in ("bullet_list_open", "ordered_list_open"):
            list_depth += 1
            ordered = token.type == "ordered_list_open"
        elif token.type in ("bullet_list_close", "ordered_list_close"):
            list_depth -= 1
        elif token.type == "table_open":
            in_table = True
            rows = []
        elif token.type == "tr_open":
            row = []
        elif token.type == "tr_close":
            rows.append(row)
        elif token.type == "table_close":
            blocks.append(block("table", rows=rows))
            in_table = False
        elif token.type == "inline":
            plain = "".join(
                t.content if t.type not in ("softbreak", "hardbreak") else "\n"
                for t in (token.children or [])
                if t.type in ("text", "code_inline", "softbreak", "hardbreak", "image")
            )
            if in_table:
                row.append(plain)
            else:
                blocks.append(
                    block(
                        "heading" if heading else "list" if list_depth else "paragraph",
                        plain,
                        level=heading or list_depth,
                        ordered=ordered,
                    )
                )
        elif token.type in ("fence", "code_block"):
            blocks.append(block("paragraph", token.content))
    return {
        "blocks": blocks,
        "warnings": (
            ["Styling is standardized. Remote Markdown images are not fetched."]
            if "![" in text
            else []
        ),
    }


def validate_text(content, name):
    text = text_of(content)
    if not text.strip():
        raise ValueError("No usable document text was found.")
    # Moodle extraction is already bounded by archive and selected-XML limits.
    # Corpus size must not be constrained by a single model request's context.
    if Path(name).suffix.lower() != '.mbz' and len(text) > MAX_TEXT:
        raise ValueError(f"Extracted document text must be {MAX_TEXT:,} characters or fewer.")


def import_file(name, data):
    mbz = Path(name).suffix.lower() == ".mbz"
    if len(data) > (MAX_MBZ_UPLOAD if mbz else MAX_UPLOAD):
        raise HTTPException(413, "Moodle backups must be 1 GB or smaller." if mbz else "Files must be 20 MB or smaller.")
    suffix = Path(name).suffix.lower()
    comments = []
    try:
        if suffix in (".txt", ".md", ".markdown"):
            decoded = data.decode("utf-8-sig")
            content = (
                markdown(decoded)
                if suffix != ".txt"
                else {
                    "blocks": [
                        block("paragraph", p)
                        for p in decoded.split("\n\n")
                        if p.strip()
                    ],
                    "warnings": [],
                }
            )
        elif suffix == ".zip":
            from app.packages import import_package

            content = import_package(data)
        elif suffix == ".mbz":
            from app.moodle import import_moodle

            content = import_moodle(data)
        elif suffix == ".pdf":
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ValueError("Encrypted PDFs are not supported.")
            if len(reader.pages) > 300:
                raise ValueError("PDFs must have 300 pages or fewer.")
            pages = [page.extract_text() or "" for page in reader.pages]
            if any(not p.strip() for p in pages):
                raise ValueError(
                    "This PDF has pages without extractable text. Convert scanned pages with OCR before uploading."
                )
            content = {
                "blocks": [block("paragraph", p) for p in pages],
                "warnings": [
                    "PDF text extracted by page. Tables, images, and reading order may need checking against the original."
                ],
            }
        elif suffix == ".docx":
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 60 * 1024 * 1024:
                    raise ValueError("Expanded Word document is too large.")
            doc = WordDocument(io.BytesIO(data))
            blocks = []
            own_export = (doc.core_properties.identifier or "").startswith("curriculum-review:")
            first_paragraph = doc.paragraphs[0]._p if doc.paragraphs else None
            for element in doc.element.body:
                if element.tag == qn("w:p"):
                    p = Paragraph(element, doc)
                    style = p.style.name if p.style else ""
                    if own_export and (style.startswith("CR Export") or (element is first_paragraph and style == "Title")):
                        continue
                    kind = (
                        "heading"
                        if style.startswith("Heading")
                        else (
                            "list"
                            if "List" in style
                            or element.find(".//" + qn("w:numPr")) is not None
                            else "paragraph"
                        )
                    )
                    level = (
                        int(style.split()[-1])
                        if kind == "heading" and style.split()[-1].isdigit()
                        else 1
                    )
                    if p.text:
                        blocks.append(
                            block(kind, p.text, level=level, ordered="Number" in style)
                        )
                    for blip in element.iter(qn("a:blip")):
                        rid = blip.get(qn("r:embed"))
                        if rid and rid in doc.part.related_parts:
                            image = doc.part.related_parts[rid]
                            if image.content_type in (
                                "image/png",
                                "image/jpeg",
                                "image/gif",
                            ):
                                blocks.append(
                                    block(
                                        "image",
                                        "Imported image",
                                        data=base64.b64encode(image.blob).decode(),
                                        mime=image.content_type,
                                    )
                                )
                elif element.tag == qn("w:tbl"):
                    t = Table(element, doc)
                    blocks.append(
                        block("table", rows=[[c.text for c in r.cells] for r in t.rows])
                    )
            active, quotes, scopes = set(), {}, {}
            for element in doc.element.body.iter():
                if element.tag == qn("w:commentRangeStart"):
                    key = element.get(qn("w:id"))
                    active.add(key)
                    quotes.setdefault(key, "")
                    if own_export:
                        parent = element.getparent()
                        while parent is not None and parent.tag != qn("w:p"):
                            parent = parent.getparent()
                        if parent is not None and Paragraph(parent,doc).style.name == "CR Export General Comment":
                            scopes[key] = "general"
                elif element.tag == qn("w:commentRangeEnd"):
                    active.discard(element.get(qn("w:id")))
                elif element.tag == qn("w:t"):
                    for key in active:
                        quotes[key] += element.text or ""
            w14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"
            w15 = "{http://schemas.microsoft.com/office/word/2012/wordml}"
            extended = {}
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "word/commentsExtended.xml" in archive.namelist():
                    root = etree.fromstring(
                        archive.read("word/commentsExtended.xml"),
                        parser=etree.XMLParser(resolve_entities=False, no_network=True),
                    )
                    extended = {
                        e.get(w15 + "paraId"): {
                            "parent": e.get(w15 + "paraIdParent"),
                            "resolved": e.get(w15 + "done") == "1",
                        }
                        for e in root
                    }
            paragraph_ids = {}
            for c in doc.comments:
                for p in c._element.iter(qn("w:p")):
                    if p.get(w14 + "paraId"):
                        paragraph_ids[p.get(w14 + "paraId")] = str(c.comment_id)
            for c in doc.comments:
                paragraph = list(c._element.iter(qn("w:p")))
                metadata = (
                    extended.get(paragraph[-1].get(w14 + "paraId"), {})
                    if paragraph
                    else {}
                )
                comments.append(
                    {
                        "source_id": str(c.comment_id),
                        "parent_id": paragraph_ids.get(metadata.get("parent")),
                        "resolved": metadata.get("resolved", False),
                        "author": c.author or "Imported author",
                        "date": str(c.timestamp or ""),
                        "text": c.text,
                        "quote": "" if scopes.get(str(c.comment_id)) == "general" else quotes.get(str(c.comment_id), ""),
                        "scope": scopes.get(str(c.comment_id), "passage"),
                    }
                )
            warnings = [
                "Word layout is standardized. Check complex fields, nested tables, tracked changes, and unsupported images against the original."
            ]
            content = {"blocks": blocks, "warnings": warnings}
        else:
            raise ValueError("Upload a .docx, .pdf, .md, .txt, Moodle .mbz, or SCORM/xAPI .zip file.")
        validate_text(content, name)
        return content, comments
    except HTTPException:
        raise
    except Exception as exc:
        message = (
            str(exc)
            if isinstance(exc, (ValueError, UnicodeDecodeError))
            else "Could not read this file. Check its format and try again."
        )
        raise HTTPException(422, message) from None


def stable_ids(content, previous):
    """Retain block identities only for unique unchanged blocks; ambiguous matches stay distinct."""

    def key(b):
        return json.dumps(semantic_block(b), sort_keys=True)

    old = {}
    for b in previous.get("blocks", []):
        old.setdefault(key(b), []).append(b["id"])
    for b in content["blocks"]:
        candidates = old.get(key(b), [])
        if len(candidates) == 1:
            b["id"] = candidates.pop()
    return content


def text_units(content):
    for b in content["blocks"]:
        if b["type"] == "table":
            for i, row in enumerate(b["rows"]):
                for j, value in enumerate(row):
                    yield {"id": f"{b['id']}:{i}:{j}", "text": value}
        elif b["type"] != "image":
            yield b


def anchor(thread, content):
    if not thread.quote:
        return {"status": "general"}
    blocks = list(text_units(content))
    same = next((b for b in blocks if b["id"] == thread.block_id), None)
    if (
        same
        and thread.start is not None
        and same.get("text", "")[thread.start : thread.end] == thread.quote
    ):
        return {
            "status": "attached",
            "block_id": same["id"],
            "start": thread.start,
            "end": thread.end,
        }
    matches = []
    for b in blocks:
        for m in re.finditer(re.escape(thread.quote), b.get("text", "")):
            matches.append(
                {
                    "status": "attached",
                    "block_id": b["id"],
                    "start": m.start(),
                    "end": m.end(),
                }
            )
    return matches[0] if len(matches) == 1 else {"status": "unplaced"}


def compare(left, right):
    def key(b):
        return json.dumps(semantic_block(b), sort_keys=True)

    a, b = left["blocks"], right["blocks"]
    ka, kb = list(map(key, a)), list(map(key, b))
    changes = []
    for tag, i, j, k, l in SequenceMatcher(None, ka, kb, autojunk=False).get_opcodes():
        if tag == "equal":
            continue
        for index in range(i, j):
            changes.append(
                {
                    "change": "moved_from" if ka[index] in kb else "removed",
                    "block": a[index],
                    "position": index + 1,
                }
            )
        for index in range(k, l):
            changes.append(
                {
                    "change": "moved_to" if kb[index] in ka else "added",
                    "block": b[index],
                    "position": index + 1,
                }
            )
    return changes


def as_markdown(content):
    result = []
    for b in content["blocks"]:
        t = b["type"]
        if t == "table":
            for index, row in enumerate(b["rows"]):
                result.append(
                    "| "
                    + " | ".join(
                        c.replace("|", "\\|").replace("\n", "<br>") for c in row
                    )
                    + " |"
                )
                if index == 0:
                    result.append("| " + " | ".join("---" for _ in row) + " |")
        elif t == "image":
            result.append(
                "![Imported image](data:" + b["mime"] + ";base64," + b["data"] + ")"
            )
        else:
            result.append(
                (
                    "#" * min(b.get("level", 1), 6) + " "
                    if t == "heading"
                    else (
                        "1. "
                        if t == "list" and b.get("ordered")
                        else "- " if t == "list" else ""
                    )
                )
                + b.get("text", "")
            )
        result.append("")
    return "\n".join(result)


def export_bytes(title, revision, fmt, threads):
    if fmt == "md":
        output = f"# {title}\n\nRevision {revision.number}\n\n" + as_markdown(
            revision.content
        )
        if threads:
            output += "\n## Comments\n"
            for t in threads:
                output += (
                    "\n> "
                    + (t["quote"] or "Whole document")
                    + "\n\n"
                    + "\n\n".join(
                        m["author_name"] + ": " + m["text"] for m in t["messages"]
                    )
                    + "\n"
                )
        return output.encode()
    doc = WordDocument()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    for style in ["Title", "Subtitle", *["Heading " + str(i) for i in range(1, 10)]]:
        doc.styles[style].font.color.rgb = RGBColor(0, 0, 0)
    doc.core_properties.identifier = 'curriculum-review:' + revision.id
    for name, base in [('CR Export Metadata','Normal'),('CR Export Appendix','Heading 1'),('CR Export General Comment','Normal'),('CR Export Unplaced Comment','Normal')]:
        style=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);style.base_style=doc.styles[base]
    doc.add_heading(title, 0)
    doc.add_paragraph(f"Revision {revision.number}", style='CR Export Metadata')
    placed = set()

    def add_text(paragraph, block_id, text):
        attached = [
            t
            for t in threads
            if t["anchor"].get("block_id") == block_id
            and t["anchor"]["status"] == "attached"
        ]
        boundaries = sorted(
            {
                0,
                len(text),
                *[
                    x
                    for t in attached
                    for x in (t["anchor"]["start"], t["anchor"]["end"])
                ],
            }
        )
        runs = [
            (s, e, paragraph.add_run(text[s:e]))
            for s, e in zip(boundaries, boundaries[1:])
        ]
        for t in attached:
            selected = [
                r
                for s, e, r in runs
                if s >= t["anchor"]["start"] and e <= t["anchor"]["end"]
            ]
            if selected:
                doc.add_comment(
                    selected,
                    text="\n\n".join(
                        f"{m['author_name']} ({m['source_date'] or m['created']}): {m['text']}"
                        for m in t["messages"]
                    ),
                    author=t["messages"][0]["author_name"],
                    initials="",
                )
                placed.add(t["id"])

    for b in revision.content["blocks"]:
        if b["type"] == "image":
            blob = base64.b64decode(b["data"])
            dimensions = WordImage.from_blob(blob)
            ratio = min(5.5 / dimensions.width.inches, 7.5 / dimensions.height.inches)
            doc.add_picture(
                io.BytesIO(blob), width=Inches(dimensions.width.inches * ratio)
            )
            continue
        if b["type"] == "table":
            rows = b["rows"]
            count = max((len(r) for r in rows), default=1)
            table = doc.add_table(rows=len(rows), cols=count)
            table.style = "Table Grid"
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                border = OxmlElement("w:" + edge)
                border.set(qn("w:val"), "single")
                border.set(qn("w:sz"), "4")
                border.set(qn("w:color"), "D9D9D9")
                borders.append(border)
            table._tbl.tblPr.append(borders)
            margins = OxmlElement("w:tblCellMar")
            for side in ("top", "left", "bottom", "right"):
                margin = OxmlElement("w:" + side)
                margin.set(qn("w:w"), "100")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            table._tbl.tblPr.append(margins)
            for i, row in enumerate(rows):
                for j, value in enumerate(row):
                    cell = table.cell(i, j)
                    add_text(cell.paragraphs[0], f"{b['id']}:{i}:{j}", value)
                    if i == 0:
                        fill = OxmlElement("w:shd")
                        fill.set(qn("w:fill"), "E8EDF1")
                        cell._tc.get_or_add_tcPr().append(fill)
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            if rows:
                repeat = OxmlElement("w:tblHeader")
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
            doc.add_paragraph()
            continue
        paragraph = doc.add_paragraph(
            style=(
                "Heading " + str(min(b.get("level", 1), 9))
                if b["type"] == "heading"
                else (
                    "List Number"
                    if b["type"] == "list" and b.get("ordered")
                    else "List Bullet" if b["type"] == "list" else "Normal"
                )
            )
        )
        add_text(paragraph, b["id"], b.get("text", ""))
    for label, items in [
        ("General comments", [t for t in threads if not t["quote"]]),
        (
            "Unplaced comments",
            [t for t in threads if t["quote"] and t["id"] not in placed],
        ),
    ]:
        if not items:
            continue
        doc.add_paragraph(label, style='CR Export Appendix')
        for t in items:
            p = doc.add_paragraph(t["quote"] or "Whole document feedback", style='CR Export General Comment' if not t['quote'] else 'CR Export Unplaced Comment')
            doc.add_comment(
                p.runs,
                text="\n\n".join(
                    m["author_name"] + ": " + m["text"] for m in t["messages"]
                ),
                author=t["messages"][0]["author_name"],
                initials="",
            )
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()
