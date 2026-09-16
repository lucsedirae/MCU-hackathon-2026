"""Integration tests use a disposable PostgreSQL database, never application tables."""

import io
import os
import unittest
import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor
from uuid import uuid4
from unittest.mock import AsyncMock, patch
import psycopg
from psycopg import sql
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from fastapi.testclient import TestClient
from docx import Document as WordDocument
from app.main import app
from app import auth, documents
from app.database import Base, database_url
from app.models import User, Revision, ReviewRun, Document, Export
from app.content import import_file, markdown, anchor, compare, fingerprint


class DocumentTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.name = "cr_test_" + uuid4().hex
        cls.connection = psycopg.connect(
            host=database_url.host,
            user=database_url.username,
            password=database_url.password,
            dbname=database_url.database,
            autocommit=True,
        )
        cls.connection.execute(
            sql.SQL("CREATE DATABASE {}").format(sql.Identifier(cls.name))
        )
        cls.engine = create_engine(database_url.set(database=cls.name))
        subprocess.run(
            [sys.executable, "-m", "alembic", "upgrade", "head"],
            env={**os.environ, "POSTGRES_DB": cls.name},
            check=True,
            capture_output=True,
        )

    @classmethod
    def tearDownClass(cls):
        cls.engine.dispose()
        cls.connection.execute(
            sql.SQL("DROP DATABASE {} WITH (FORCE)").format(sql.Identifier(cls.name))
        )
        cls.connection.close()

    def setUp(self):
        with self.engine.begin() as c:
            for table in reversed(Base.metadata.sorted_tables):
                c.execute(table.delete())
        auth._attempts.clear()

        def session():
            with Session(self.engine) as s:
                yield s

        app.dependency_overrides[auth.db] = session
        self.engine_patch = patch.object(documents, "engine", self.engine)
        self.engine_patch.start()
        self.admin = TestClient(app)
        self.owner = TestClient(app)
        self.member = TestClient(app)
        self.assertEqual(
            self.admin.post(
                "/api/auth/setup",
                json={
                    "email": "admin@test.local",
                    "name": "Admin",
                    "password": "administrator-secret",
                },
            ).status_code,
            200,
        )
        self.admin.post(
            "/api/auth/login",
            json={"email": "admin@test.local", "password": "administrator-secret"},
        )
        self.ids = {}
        for client, name in [(self.owner, "Owner"), (self.member, "Member")]:
            r = self.admin.post(
                "/api/auth/users",
                json={
                    "email": name.lower() + "@test.local",
                    "name": name,
                    "password": "temporary-secret",
                },
            )
            self.ids[name] = r.json()["id"]
            client.post(
                "/api/auth/login",
                json={
                    "email": name.lower() + "@test.local",
                    "password": "temporary-secret",
                },
            )
            self.assertEqual(client.get("/api/workspaces").status_code, 403)
            client.post(
                "/api/auth/password",
                json={
                    "old_password": "temporary-secret",
                    "new_password": "changed-password",
                },
            )
            client.post(
                "/api/auth/login",
                json={
                    "email": name.lower() + "@test.local",
                    "password": "changed-password",
                },
            )
        w = self.admin.post(
            "/api/workspaces",
            json={"title": "Curriculum testing", "owner_id": self.ids["Owner"]},
        )
        self.wid = w.json()["id"]
        self.did = self.owner.get("/api/workspaces/" + self.wid).json()["documents"][0][
            "id"
        ]

    def tearDown(self):
        for c in (self.admin, self.owner, self.member):
            c.close()
        self.engine_patch.stop()
        app.dependency_overrides.clear()

    def upload(
        self,
        text="# Curriculum\n\nLearning objectives\n\n- Practice safely",
        base="",
        current="",
        client=None,
        name="curriculum.md",
    ):
        r = (client or self.owner).post(
            f"/api/documents/{self.did}/upload",
            files={"file": (name, text.encode())},
            data={"base_id": base, "expected_current": current},
        )
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()["id"]

    def accept(self, rid, current=None, resolve=None):
        prefix = f"/api/documents/{self.did}/revisions/{rid}"
        cmp = self.owner.post(prefix + "/compare", json={}).json()
        response = self.owner.post(
            prefix + "/accept",
            json={
                "comparison_token": cmp["token"],
                "expected_current": current,
                "resolve_threads": resolve or [],
            },
        )
        self.assertEqual(response.status_code, 200, response.text)

    def test_permissions_auth_and_setup(self):
        self.assertEqual(TestClient(app).get("/api/workspaces").status_code, 401)
        self.assertEqual(
            self.admin.post(
                "/api/auth/setup",
                json={
                    "email": "x@test.local",
                    "name": "X",
                    "password": "long-enough-secret",
                },
            ).status_code,
            409,
        )
        self.assertEqual(
            self.member.post(
                "/api/workspaces", json={"title": "No", "owner_id": self.ids["Member"]}
            ).status_code,
            403,
        )
        for client in (self.admin, self.member):
            self.assertEqual(
                client.post(
                    f"/api/documents/{self.did}/upload",
                    files={"file": ("x.txt", b"No")},
                ).status_code,
                403,
            )
            self.assertEqual(
                client.post(
                    f"/api/workspaces/{self.wid}/runs",
                    json={"kind": "generate", "instructions": "Hello"},
                ).status_code,
                403,
            )
        self.assertEqual(self.member.get("/api/settings/llm").status_code, 403)
        self.assertEqual(
            self.owner.post(
                "/api/auth/logout", headers={"sec-fetch-site": "cross-site"}
            ).status_code,
            403,
        )

    def test_proposals_stale_acceptance_restore_and_dedupe(self):
        one = self.upload()
        self.accept(one)
        two = self.upload("# Curriculum\n\nChanged objectives", one, one)
        three = self.upload("# Curriculum\n\nAlternative objectives", one, one)
        prefix = f"/api/documents/{self.did}/revisions/{three}"
        token = self.owner.post(prefix + "/compare", json={}).json()["token"]
        self.accept(two, one)
        stale = self.owner.post(
            prefix + "/accept",
            json={"comparison_token": token, "expected_current": one},
        )
        self.assertEqual(stale.status_code, 409)
        self.accept(three, two)
        restored = self.owner.post(
            f"/api/documents/{self.did}/revisions/{one}/restore"
        ).json()
        self.assertEqual(restored["status"], "pending")
        self.assertEqual(
            self.owner.get(f"/api/documents/{self.did}").json()["current_id"], three
        )
        reject = f'/api/documents/{self.did}/revisions/{restored["id"]}/state'
        self.assertEqual(
            self.owner.post(reject, json={"action": "reject"}).json()["status"],
            "rejected",
        )
        self.assertEqual(
            self.owner.post(reject, json={"action": "reopen"}).json()["status"],
            "pending",
        )
        same = self.upload("# Curriculum\n\nAlternative objectives", three, three)
        self.assertEqual(same, three)
        self.assertEqual(
            self.owner.get(f"/api/documents/{self.did}/revisions/{one}").json()[
                "content"
            ]["blocks"][1]["text"],
            "Learning objectives",
        )

    def test_comments_carry_forward_export_and_import(self):
        one = self.upload()
        self.accept(one)
        prefix = f"/api/documents/{self.did}/revisions/"
        content = self.owner.get(prefix + one).json()["content"]
        b = content["blocks"][1]
        t = self.member.post(
            prefix + one + "/comments",
            json={"text": "Please clarify", "block_id": b["id"], "start": 0, "end": 8},
        ).json()["id"]
        self.assertEqual(
            self.admin.put(
                "/api/comments/" + t + "/state", json={"resolved": True}
            ).status_code,
            403,
        )
        self.assertEqual(
            self.member.post(
                "/api/comments/" + t + "/replies", json={"text": "Include examples"}
            ).status_code,
            200,
        )
        two = self.upload(
            "# Curriculum\n\nLearning objectives\n\nNew section", one, one
        )
        rev = self.owner.get(prefix + two).json()
        self.assertEqual(rev["threads"][0]["anchor"]["status"], "attached")
        first = self.owner.get(prefix + two + "/export?format=docx&comments=true")
        self.assertEqual(
            first.status_code, 200, first.text[:100] if first.status_code != 200 else ""
        )
        self.assertEqual(
            first.content,
            self.owner.get(prefix + two + "/export?format=docx&comments=true").content,
        )
        first_export_id = self.owner.get(prefix + two).json()["exports"][0]["id"]
        doc = WordDocument(io.BytesIO(first.content))
        self.assertEqual(len(doc.comments), 1)
        self.assertIn("Include examples", doc.comments.get(0).text)
        imported, comments = import_file("roundtrip.docx", first.content)
        self.assertEqual(comments[0]["quote"], "Learning")
        self.assertEqual(fingerprint(imported), fingerprint(rev["content"]))
        self.assertIn("Please clarify", comments[0]["text"])
        self.member.post(
            "/api/comments/" + t + "/replies", json={"text": "More context"}
        )
        second = self.owner.get(prefix + two + "/export?format=docx&comments=true")
        self.assertNotEqual(first.content, second.content)
        self.assertEqual(self.owner.get("/api/exports/" + first_export_id).content, first.content)
        with Session(self.engine) as s:
            self.assertEqual(len(list(s.scalars(select(Export)))), 2)
        roundtrip = self.owner.post(
            f'/api/documents/{self.did}/upload',
            files={'file':('roundtrip.docx',first.content)},
            data={'base_id':two,'expected_current':one},
        )
        self.assertEqual(roundtrip.status_code,200,roundtrip.text)
        roundtrip_threads = self.owner.get(prefix+roundtrip.json()['id']).json()['threads']
        self.assertEqual(len(roundtrip_threads),1)
        self.assertEqual(roundtrip_threads[0]['id'],t)
        three = self.upload("# Curriculum\n\nRemoved everything", two, one)
        self.assertEqual(
            self.owner.get(prefix + three).json()["threads"][0]["anchor"]["status"],
            "unplaced",
        )
        self.assertEqual(
            self.member.put(
                "/api/comments/" + t + "/state", json={"resolved": True}
            ).status_code,
            200,
        )

    def test_runs_reports_transcripts_and_generation(self):
        one = self.upload()
        mock = AsyncMock(
            return_value={
                "text": "# Review findings\n\nClear objectives.",
                "model": "mock-model",
            }
        )
        with patch.object(documents, "request_completion", mock):
            for _ in range(2):
                response = self.owner.post(
                    f"/api/workspaces/{self.wid}/runs",
                    json={"revision_id": one, "kind": "review"},
                )
                self.assertEqual(response.status_code, 200, response.text)
        w = self.owner.get("/api/workspaces/" + self.wid).json()
        self.assertEqual(len([d for d in w["documents"] if d["kind"] == "report"]), 2)
        self.assertEqual(
            len([d for d in w["documents"] if d["kind"] == "transcript"]), 2
        )
        self.assertTrue(
            all(r["was_proposal"] and r["status"] == "completed" for r in w["runs"])
        )
        transcript = next(d for d in w["documents"] if d["kind"] == "transcript")
        self.assertEqual(
            self.owner.post(
                f'/api/documents/{transcript["id"]}/upload',
                files={"file": ("x.txt", b"changed")},
            ).status_code,
            403,
        )
        report = next(d for d in w["documents"] if d["kind"] == "report")
        response = self.owner.post(
            f'/api/documents/{report["id"]}/upload',
            files={"file": ("report.txt", b"Edited report")},
            data={"expected_current": report["current_id"]},
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(
            self.owner.get(f'/api/documents/{report["id"]}').json()["current_id"],
            response.json()["id"],
        )
        with patch.object(
            documents,
            "request_completion",
            AsyncMock(
                return_value={
                    "text": "# New curriculum\n\nNew objectives",
                    "model": "mock",
                }
            ),
        ):
            self.owner.post(
                f"/api/workspaces/{self.wid}/runs",
                json={"revision_id": one, "kind": "generate"},
            )
        d = self.owner.get("/api/documents/" + self.did).json()
        self.assertIsNone(d["current_id"])
        self.assertEqual(d["revisions"][0]["status"], "pending")
        bad = self.owner.get(
            f'/api/documents/{self.did}/revisions/{report["current_id"]}'
        )
        self.assertEqual(bad.status_code, 404)

    def test_parallel_revision_numbers_and_failed_run(self):
        one = self.upload()
        self.accept(one)

        def upload(number):
            client = TestClient(app)
            client.cookies.update(self.owner.cookies)
            return client.post(
                f"/api/documents/{self.did}/upload",
                files={"file": ("x.txt", f"Proposal {number}".encode())},
                data={"base_id": one, "expected_current": one},
            )

        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(upload, [1, 2]))
        self.assertTrue(
            all(r.status_code == 200 for r in results), [r.text for r in results]
        )
        self.assertEqual({r.json()["number"] for r in results}, {2, 3})
        with patch.object(
            documents,
            "request_completion",
            AsyncMock(side_effect=RuntimeError("private error")),
        ):
            self.owner.post(
                f"/api/workspaces/{self.wid}/runs",
                json={"revision_id": one, "kind": "review"},
            )
        run = self.owner.get("/api/workspaces/" + self.wid).json()["runs"][0]
        self.assertEqual(run["status"], "failed")
        self.assertIsNotNone(run["transcript_id"])
        self.assertNotIn("private error", str(run))

    def test_table_comments_and_alternative_branch_history(self):
        one = self.upload(
            "# Curriculum\n\n| Skill | Evidence |\n| --- | --- |\n| Assessment | Practice |"
        )
        self.accept(one)
        prefix = f"/api/documents/{self.did}/revisions/"
        table = self.owner.get(prefix + one).json()["content"]["blocks"][1]
        thread = self.member.post(
            prefix + one + "/comments",
            json={
                "text": "Add a checklist",
                "block_id": table["id"] + ":1:1",
                "start": 0,
                "end": 8,
            },
        ).json()["id"]
        exported = self.owner.get(prefix + one + "/export?format=docx&comments=true")
        content, comments = import_file("table.docx", exported.content)
        self.assertEqual(comments[0]["quote"], "Practice")
        two = self.upload("Alternative two", one, one)
        three = self.upload("Alternative three", one, one)
        self.accept(two, one)
        new_thread = self.member.post(
            prefix + two + "/comments",
            json={"text": "Feedback from the accepted alternative"},
        ).json()["id"]
        self.accept(three, two)
        ids = {t["id"] for t in self.owner.get(prefix + three).json()["threads"]}
        self.assertIn(thread, ids)
        self.assertIn(new_thread, ids)

    def test_ownership_transfer_reset_and_interruption(self):
        one = self.upload()
        response = self.admin.put(
            "/api/workspaces/" + self.wid + "/owner",
            json={"owner_id": self.ids["Member"]},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            self.owner.post(
                f"/api/documents/{self.did}/revisions/{one}/restore"
            ).status_code,
            403,
        )
        self.assertEqual(
            self.member.post(
                f"/api/documents/{self.did}/revisions/{one}/restore"
            ).status_code,
            200,
        )
        with Session(self.engine) as session:
            run = ReviewRun(
                workspace_id=self.wid,
                revision_id=one,
                author_id=self.ids["Member"],
                kind="review",
                entries=[{"timestamp": "2026-09-15T12:00:00Z", "text": "Started"}],
            )
            session.add(run)
            session.commit()
            run_id = run.id
        documents.recover_runs()
        with Session(self.engine) as session:
            run = session.get(ReviewRun, run_id)
            self.assertEqual(run.status, "interrupted")
            self.assertIsNotNone(run.transcript_id)
        reset = self.admin.post(
            "/api/auth/users/" + self.ids["Member"] + "/reset",
            json={"password": "replacement-temporary"},
        )
        self.assertEqual(reset.status_code, 200)
        self.assertEqual(self.member.get("/api/workspaces").status_code, 401)

    def test_moodle_upload_proposal_original_and_word_export(self):
        from test_moodle import backup_bytes
        data = backup_bytes('gzip')
        response = self.owner.post(
            f'/api/documents/{self.did}/upload',
            files={'file': ('course.mbz', data, 'application/octet-stream')},
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['status'], 'pending')
        rid = response.json()['id']
        revision = self.owner.get(f'/api/documents/{self.did}/revisions/{rid}').json()
        self.assertEqual(revision['content']['blocks'][0]['text'], 'Clinical skills')
        original = self.owner.get('/api/originals/' + revision['originals'][0]['id'])
        self.assertEqual(original.content, data)
        self.accept(rid)
        exported = self.owner.get(f'/api/documents/{self.did}/revisions/{rid}/export?format=docx')
        self.assertEqual(exported.status_code, 200)
        self.assertIn('Clinical skills', '\n'.join(p.text for p in WordDocument(io.BytesIO(exported.content)).paragraphs))

    def test_formats_and_structure(self):
        doc = WordDocument()
        doc.add_heading("Objectives", 1)
        doc.add_paragraph("Learn safely", style="List Bullet")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Skill"
        table.cell(0, 1).text = "Evidence"
        stream = io.BytesIO()
        doc.save(stream)
        content, _ = import_file("example.docx", stream.getvalue())
        self.assertEqual(
            [b["type"] for b in content["blocks"]], ["heading", "list", "table"]
        )
        md = markdown("# Title\n\n| Skill | Evidence |\n| --- | --- |\n| A | B |")
        self.assertEqual(md["blocks"][1]["type"], "table")
        reordered = {"blocks": list(reversed(content["blocks"]))}
        self.assertTrue(
            any(c["change"].startswith("moved") for c in compare(content, reordered))
        )
        from pypdf import PdfWriter

        pdf = PdfWriter()
        pdf.add_blank_page(width=200, height=200)
        output = io.BytesIO()
        pdf.write(output)
        from fastapi import HTTPException

        with self.assertRaises(HTTPException) as ctx:
            import_file("scan.pdf", output.getvalue())
        self.assertIn("OCR", ctx.exception.detail)


if __name__ == "__main__":
    unittest.main()
